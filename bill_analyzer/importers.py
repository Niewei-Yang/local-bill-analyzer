from __future__ import annotations

import csv
import io
from abc import ABC, abstractmethod
from pathlib import Path

from openpyxl import load_workbook

from .models import Transaction, clean_text, parse_amount_cents, parse_datetime


class UnsupportedBillError(ValueError):
    pass


class BillImporter(ABC):
    platform: str

    @abstractmethod
    def score(self, filename: str, content: bytes) -> int:
        raise NotImplementedError

    @abstractmethod
    def parse(self, filename: str, content: bytes) -> list[Transaction]:
        raise NotImplementedError


def _direction(value: object) -> str:
    return {"支出": "expense", "收入": "income", "不计收支": "neutral"}.get(clean_text(value), "neutral")


class WeChatImporter(BillImporter):
    platform = "微信"

    def score(self, filename: str, content: bytes) -> int:
        name = filename.lower()
        if name.endswith(".xlsx") and ("微信" in filename or content[:2] == b"PK"):
            return 80 + (20 if "微信" in filename else 0)
        return 0

    def parse(self, filename: str, content: bytes) -> list[Transaction]:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [tuple(cell.value for cell in row) for row in sheet.iter_rows()]
        workbook.close()
        header_index = next(
            (i for i, row in enumerate(rows[:40]) if "交易时间" in row and "收/支" in row and "金额(元)" in row),
            None,
        )
        if header_index is None:
            raise UnsupportedBillError("未找到微信账单表头")
        headers = [clean_text(value) for value in rows[header_index]]
        records: list[Transaction] = []
        for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            values = dict(zip(headers, row))
            if not clean_text(values.get("交易时间")):
                continue
            try:
                direction = _direction(values.get("收/支"))
                tx_type = clean_text(values.get("交易类型"))
                status = clean_text(values.get("当前状态"))
                is_refund = int(direction != "expense" and ("退款" in tx_type or "退款" in status))
                records.append(
                    Transaction(
                        platform=self.platform,
                        transaction_time=parse_datetime(values.get("交易时间")),
                        direction=direction,
                        amount_cents=parse_amount_cents(values.get("金额(元)")),
                        category_raw=tx_type,
                        counterparty=clean_text(values.get("交易对方")),
                        description=clean_text(values.get("商品")),
                        payment_method=clean_text(values.get("支付方式")),
                        status=status,
                        source_order_id=clean_text(values.get("交易单号")),
                        merchant_order_id=clean_text(values.get("商户单号")),
                        note=clean_text(values.get("备注")),
                        is_refund=is_refund,
                    )
                )
            except ValueError as exc:
                raise ValueError(f"{filename} 第 {row_number} 行：{exc}") from exc
        return records


class AlipayImporter(BillImporter):
    platform = "支付宝"

    def score(self, filename: str, content: bytes) -> int:
        if not filename.lower().endswith(".csv"):
            return 0
        bonus = 40 if "支付宝" in filename else 0
        head = content[:3000]
        for encoding in ("gb18030", "utf-8-sig", "utf-8"):
            try:
                if "支付宝" in head.decode(encoding):
                    return 60 + bonus
            except UnicodeDecodeError:
                continue
        return 30 + bonus

    def parse(self, filename: str, content: bytes) -> list[Transaction]:
        text = None
        for encoding in ("gb18030", "utf-8-sig", "utf-8"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise UnsupportedBillError("支付宝 CSV 编码无法识别")
        rows = list(csv.reader(io.StringIO(text)))
        header_index = next(
            (i for i, row in enumerate(rows[:60]) if "交易时间" in row and "收/支" in row and "金额" in row),
            None,
        )
        if header_index is None:
            raise UnsupportedBillError("未找到支付宝账单表头")
        headers = [clean_text(value) for value in rows[header_index]]
        records: list[Transaction] = []
        for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if not row or not clean_text(row[0]):
                continue
            values = dict(zip(headers, row))
            try:
                direction = _direction(values.get("收/支"))
                status = clean_text(values.get("交易状态"))
                category_raw = clean_text(values.get("交易分类"))
                is_refund = int(direction != "expense" and ("退款" in status or category_raw == "退款"))
                records.append(
                    Transaction(
                        platform=self.platform,
                        transaction_time=parse_datetime(values.get("交易时间")),
                        direction=direction,
                        amount_cents=parse_amount_cents(values.get("金额")),
                        category_raw=category_raw,
                        counterparty=clean_text(values.get("交易对方")),
                        description=clean_text(values.get("商品说明")),
                        payment_method=clean_text(values.get("收/付款方式")),
                        status=status,
                        source_order_id=clean_text(values.get("交易订单号")),
                        merchant_order_id=clean_text(values.get("商家订单号")),
                        note=clean_text(values.get("备注")),
                        is_refund=is_refund,
                    )
                )
            except ValueError as exc:
                raise ValueError(f"{filename} 第 {row_number} 行：{exc}") from exc
        return records


IMPORTERS: list[BillImporter] = [WeChatImporter(), AlipayImporter()]


def choose_importer(filename: str, content: bytes) -> BillImporter:
    ranked = sorted(((item.score(filename, content), item) for item in IMPORTERS), reverse=True, key=lambda x: x[0])
    if not ranked or ranked[0][0] <= 0:
        suffix = Path(filename).suffix or "未知格式"
        raise UnsupportedBillError(f"暂不支持 {suffix} 文件")
    return ranked[0][1]
